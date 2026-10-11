"""AI Resume and Job Matcher - Flask application."""

from __future__ import annotations

import io
import logging
import os
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template, request, url_for
from pypdf import PdfReader
from werkzeug.utils import secure_filename

from matcher import analyze_match

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY", "local-development-key-change-me"),
    MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES,
)
logging.basicConfig(level=os.environ.get("LOG_LEVEL", "INFO"))


def extract_resume_text(upload) -> str:
    """Extract text from a supported upload without saving it to disk."""
    filename = secure_filename(upload.filename or "")
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Upload a PDF, DOCX, or TXT file.")
    raw = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("The uploaded file is larger than 5 MB.")
    if not raw:
        raise ValueError("The selected file is empty.")

    try:
        if suffix == ".txt":
            return raw.decode("utf-8-sig")
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        document = Document(io.BytesIO(raw))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(paragraphs)
    except ValueError:
        raise
    except Exception as exc:
        app.logger.info("Resume parsing failed for %s: %s", suffix, exc)
        raise ValueError("We could not read that file. Try another file or paste the resume text.") from exc


@app.get("/")
def index():
    return render_template("index.html", result=None)


@app.post("/match")
def match():
    resume_text = request.form.get("resume_text", "").strip()
    job_text = request.form.get("job_text", "").strip()
    upload = request.files.get("resume_file")

    try:
        if upload and upload.filename:
            resume_text = extract_resume_text(upload).strip()
        if not resume_text:
            raise ValueError("Paste resume text or choose a resume file.")
        if not job_text:
            raise ValueError("Paste a job description to compare with the resume.")
        result = analyze_match(resume_text, job_text)
        return render_template("index.html", result=result,
                               resume_text=resume_text, job_text=job_text)
    except ValueError as exc:
        flash(str(exc), "error")
        return render_template("index.html", result=None,
                               resume_text=resume_text, job_text=job_text), 400


@app.errorhandler(413)
def file_too_large(_error):
    flash("The uploaded file is larger than 5 MB.", "error")
    return render_template("index.html", result=None), 413


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
